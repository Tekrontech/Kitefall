# Builds docs/Kitefall_3D_Characters.xlsx: everything an artist (or an AI helper) needs to make the 3D character models
# that replace the code-built characters. Data comes from Assets/Scripts/Core/Meta/CharacterCatalog.cs.
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/Kitefall_3D_Characters.xlsx"
src = open("Assets/Scripts/Core/Meta/CharacterCatalog.cs").read()
chars = re.findall(r'Char\("(\w+)", "([^"]+)", "([^"]+)", "(#\w+)", "(#\w+)", (\d+), ([\d.]+)f, ([\d.]+)f, (\d+), (\d+), "(\w+)"\)', src)
outfits = {o[0]: o for o in re.findall(r'Outfit\("(\w+)", "([^"]+)", Rarity\.\w+, (\d), "(#\w+)", "(#\w+)"', src)}
HAIR = {0: "short crop", 1: "side part, a bit longer", 2: "curly", 3: "long straight", 4: "bun", 5: "bob", 6: "short hair + short beard", 7: "long ponytail"}
STYLE = {0: "long kurta tunic + loose trousers (shalwar)", 1: "open jacket over shirt + trousers", 2: "hoodie / track top + trousers",
         3: "waistcoat or sherwani over long tunic", 4: "festival tunic with sash and embroidered bands", 5: "short-sleeve sports / beach shirt + trousers"}
LOOK = {
    "rookie_m": "cheerful young man, warm brown skin, short dark hair", "rookie_f": "confident young woman, long dark hair",
    "zara": "stylish young woman from Lahore, hair in a high bun, playful smile", "rafay": "tall broad-shouldered young man from Karachi, buzz cut",
    "lucas": "friendly young man from Lisbon, wavy chestnut hair, fair skin", "aiko": "energetic young woman from Osaka, sleek black bob",
    "niran": "calm young man from Chiang Mai, short black hair", "farid": "rugged young man from Kabul, tousled dark hair, short beard",
    "siti": "quick athletic young woman from Kuala Lumpur, long ponytail", "arjun": "dignified young man from Ahmedabad, neat black hair",
    "mateus": "cheerful young man from Rio de Janeiro, short curly hair", "valentina": "confident young woman from Santiago, long wavy brown hair",
    "diego": "friendly young man from Guatemala City, short black hair", "emma": "bright young woman from coastal France, honey-blonde bob",
    "jack": "easy-going young man from California, brown side-parted hair", "amara": "joyful young woman from Lagos, high bun, deep brown skin",
    "youssef": "calm young man from Marrakesh, short dark hair, neat beard", "elif": "lively young woman from Istanbul, long dark ponytail",
    "nikos": "sunny young man from Athens, short dark hair, olive skin"}
COMMON = ("stylised semi-realistic mobile game character, friendly proportions (slightly larger head and hands), clean readable shapes, "
          "hand-painted look textures, full body, standing in A-pose (arms 45° down, legs slightly apart), facing the camera, plain background, "
          "no weapons, no text, no logos, original character (not a real person), adult 18+ look")

wb = Workbook()
H = PatternFill("solid", fgColor="3A2A6B"); HF = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D9CFB8"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def sheet(ws, headers, widths, rows, height=60):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, i, h); c.fill = H; c.font = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v); c.alignment = wrap; c.border = B
        ws.row_dimensions[r].height = height
    ws.freeze_panes = "C2"

# ---- 1. Steps
ws = wb.active; ws.title = "1 Steps (padhein)"
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 130
ws["B1"] = "Kitefall – 3D character models: poora tareeqa"; ws["B1"].font = Font(bold=True, size=16, color="3A2A6B")
steps = [
    ("Maqsad", "Game ke code-wale simple characters ki jagah aap ke portraits jaise asli 3D models. Har character ka ek model (.fbx)."),
    ("1", "Character chunein (sheet '2 Characters'). Pehle 1 character banayein (jaise Sana), game mein test karein, phir baaqi."),
    ("2", "Full-body tasveer: Excel 'char_<naam>_fullbody' wali tasveer, ya sheet 2 ka 'Image prompt' se nayi banayein – A-pose, seedha khada, saada background."),
    ("3", "Image → 3D: Meshy.ai ya Tripo3d.ai par 'Image to 3D' → tasveer upload → Settings sheet '4 Technical spec' ke mutabiq (poly / texture) → FBX download."),
    ("4", "Rig: Mixamo.com (free) → Upload Character → markers lagayein (chin, wrists, elbows, knees, groin) → Next → 'Download' FBX (T-pose ya A-pose, with skin)."),
    ("5", "Animations: Mixamo se sheet '3 Animations' wali clips download karein – Format: FBX for Unity, Skin: Without Skin, 30 FPS."),
    ("6", "File naam: model = <character id>.fbx (sheet 2 'Model file name'), animations = sheet 3 wale naam. Sab Unity mein Assets/Art/Characters/ folder mein."),
    ("7", "Unity: Kitefall ▸ Characters ▸ Setup Models (Claude Code yeh menu banayega) – model Humanoid set hoga, animations jud jayengi, game khud use karega."),
    ("8", "Check: Wardrobe mein character Equip → Home + match mein naya model. Size / haath galat lage to screenshot bhejein."),
    ("Claude se", "Claude (chat) khud .fbx file nahi bana sakta. Sheet '5 Prompts for Claude' mein do kaam ke prompts hain: (a) Meshy/Tripo ke liye behtar prompt likhwana, (b) Blender script se simple stylised model banwana (agar AI 3D tool use na karna ho)."),
    ("Outfits", "Model apne default outfit mein banega (sheet 2). Doosre outfits ke liye alag model ban sakta hai: <character id>__<outfit id>.fbx (optional, baad mein)."),
    ("Safety", "Asli logon jaisa chehra, brand logos, ya kisi aur game ke character mat banayein. Dor (string) cotton/silk jaisi – glass ya metal nahi."),
]
for i, (a, b) in enumerate(steps, 3):
    ws.cell(i, 1, a).font = Font(bold=True); c = ws.cell(i, 2, b); c.alignment = wrap
    ws.row_dimensions[i].height = 34

# ---- 2. Characters
rows = []
for n, (cid, name, region, skin, hair, hs, height, build, lvl, gems, outfit) in enumerate(chars, 1):
    o = outfits.get(outfit)
    oname, ostyle, p1, p2 = (o[1], int(o[2]), o[3], o[4]) if o else (outfit, 0, "", "")
    h_m = round(1.75 * float(height), 2)
    desc = f"{LOOK.get(cid, name)}; skin {skin}, hair {hair} ({HAIR[int(hs)]}); outfit '{oname}': {STYLE[ostyle]}, main colour {p1}, second colour {p2}"
    img_prompt = f"Full-body character design of {name}, {LOOK.get(cid, '')}, wearing {STYLE[ostyle]} in {p1} and {p2}, {COMMON}"
    to3d = f"{name}: {LOOK.get(cid, '')}, {STYLE[ostyle]}, colours {p1} / {p2}; A-pose, game-ready, quad-friendly topology, PBR textures, ~{12000 if cid.startswith('rookie') else 10000} triangles"
    rows.append([n, cid, name, region, f"{cid}.fbx", f"{h_m} m", f"{float(build):.2f}", skin, hair, HAIR[int(hs)], f"{oname} ({outfit})", STYLE[ostyle], p1, p2,
                 "free" if int(gems) == 0 and int(lvl) <= 1 else f"level {lvl}" if int(gems) == 0 else f"{gems} gems",
                 desc, img_prompt, to3d, f"char_{cid}_fullbody.png / char_{cid}_portrait.png", "1" if cid in ("rookie_m", "rookie_f") else "2", ""])
ws = wb.create_sheet("2 Characters")
sheet(ws, ["#", "Character id", "Name", "Home town", "Model file name", "Height", "Build", "Skin", "Hair colour", "Hair style", "Default outfit", "Outfit shape",
           "Outfit colour 1", "Outfit colour 2", "Unlock", "Look description", "Image prompt (full-body A-pose)", "Meshy / Tripo text (image-to-3D notes)",
           "Reference images already made", "Priority", "Done?"],
          [5, 13, 11, 16, 16, 9, 7, 9, 10, 18, 26, 30, 10, 10, 10, 50, 70, 60, 34, 8, 8], rows, 120)
for r in range(2, len(rows) + 2):
    for col in (8, 9, 13, 14):
        v = ws.cell(r, col).value
        if v and v.startswith("#"): ws.cell(r, col).fill = PatternFill("solid", fgColor=v[1:])

# ---- 3. Animations
anims = [
    ("idle", "Idle (any calm standing idle)", "Standing on the roof, between rounds, wardrobe", "Loop", "Must"),
    ("walk", "Walking", "Moving on the roof, roof hopping", "Loop, In Place ON", "Must"),
    ("run", "Running", "Fast moves, warmup plaza", "Loop, In Place ON", "Must"),
    ("fly_hold", "Standing Idle / 'Holding' pose with arms raised a little", "Flying the kite (the game moves the hands onto line + spool with IK)", "Loop", "Must"),
    ("jump", "Jump", "Jump links between roofs", "Once", "Nice"),
    ("climb", "Climbing Ladder", "Ladder links", "Loop", "Nice"),
    ("victory", "Victory / Cheering", "Win emote, results", "Once", "Nice"),
    ("clap", "Clapping", "Clap emote", "Once", "Nice"),
    ("wave", "Waving", "Wave emote", "Once", "Nice"),
    ("dance", "Dancing (any happy dance)", "Victory dance emote", "Loop", "Nice"),
    ("sad", "Defeated / Disappointed", "Kite got cut", "Once", "Nice"),
]
ws = wb.create_sheet("3 Animations")
sheet(ws, ["File name (Assets/Art/Characters/Anim/)", "Mixamo search", "Used for", "Settings", "Need"], [36, 46, 52, 22, 8],
      [[f"anim_{a}.fbx", b, c, d, e] for a, b, c, d, e in anims], 36)
ws.cell(len(anims) + 3, 1, "One set of animations works for ALL characters (Unity Humanoid retargeting). Download with: FBX for Unity, Without Skin, 30 fps.").font = Font(bold=True)

# ---- 4. Technical spec
spec = [
    ("Format", ".fbx (binary). Y-up. 1 unit = 1 metre."),
    ("Size", "Height as in sheet 2 (≈ 1.75 m × height factor). Feet on the ground at (0,0,0)."),
    ("Facing", "Character looks toward +Z (Unity forward). Mixamo default is fine."),
    ("Pose", "A-pose or T-pose in the rig file."),
    ("Rig", "Mixamo auto-rig (standard 'mixamorig' bone names). Unity will use it as Humanoid."),
    ("Triangles", "8,000 – 15,000 per character (mobile; up to 12 characters in one match)."),
    ("Textures", "One material if possible. Base colour 1024×1024 (2048 max), optional normal map. PNG/JPG."),
    ("Hands", "Open hands, fingers separated a little (the game puts the line in the left hand and the spool in the right)."),
    ("Head", "Neutral, friendly face. Hair as one solid mesh (no hair cards needed)."),
    ("Accessories", "Do NOT model caps/glasses/bags – the game adds accessories itself on head, eyes, back and waist."),
    ("Folder", "Assets/Art/Characters/<id>.fbx  and  Assets/Art/Characters/Anim/anim_*.fbx"),
    ("Optional outfits", "Assets/Art/Characters/<id>__<outfit id>.fbx – same rig, different clothes."),
]
ws = wb.create_sheet("4 Technical spec")
sheet(ws, ["Item", "Requirement"], [22, 120], spec, 30)

# ---- 5. Prompts for Claude
p_a = ("I am making 3D characters for my mobile kite-flying game Kitefall. For the character below, write me (1) a clean image prompt for a "
       "full-body A-pose character sheet (front view, plain background) and (2) the best short text to paste into Meshy.ai / Tripo3d.ai "
       "'Image to 3D' so the model is game-ready (8–15k triangles, one 1024px texture, A-pose, Mixamo-friendly). Stylised semi-realistic, "
       "friendly mobile game style, original character, adult. Character: <paste one row from sheet '2 Characters': name, look description, outfit, colours>")
p_b = ("Write a Blender 4.x Python script (bpy) that builds a stylised low-poly mobile-game character, rigged for Mixamo-style humanoid use, and exports "
       "it as FBX. Requirements: Y-up export, 1 unit = 1 m, height <HEIGHT> m, feet at origin, facing +Z after export, A-pose, simple separate meshes "
       "for head, hair (<HAIR STYLE>), torso, arms, hands, legs, shoes joined into one mesh, one material per colour: skin <SKIN HEX>, hair <HAIR HEX>, "
       "outfit main <COLOUR 1>, outfit second <COLOUR 2>, outfit shape: <OUTFIT SHAPE>. Create an armature with bones named like Mixamo "
       "(mixamorig:Hips, Spine, Spine1, Spine2, Neck, Head, LeftShoulder, LeftArm, LeftForeArm, LeftHand, RightShoulder, RightArm, RightForeArm, "
       "RightHand, LeftUpLeg, LeftLeg, LeftFoot, RightUpLeg, RightLeg, RightFoot) and parent the mesh with automatic weights. Keep it under "
       "10,000 triangles. Export to //<ID>.fbx with 'Apply Transform' on. Explain how to run the script in Blender.")
p_c = ("Here is my Kitefall character spec (sheets 2–4 of Kitefall_3D_Characters.xlsx). Check my exported FBX settings list and tell me what is "
       "wrong before I import it into Unity: <paste your export settings / Mixamo download settings>")
ws = wb.create_sheet("5 Prompts for Claude")
sheet(ws, ["Use", "Prompt (copy, fill the <...> parts)"], [30, 140],
      [["(a) Better image + Meshy/Tripo prompt", p_a], ["(b) Blender script – simple stylised model without AI 3D tools", p_b], ["(c) Check settings before import", p_c]], 150)
wb.save(OUT)
print("saved", OUT, len(rows), "characters")
