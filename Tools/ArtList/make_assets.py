import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

STYLE = ("classical hand-painted storybook illustration in a timeless world kite-festival style, premium mobile game art, "
         "gouache and watercolour texture on warm ivory paper, fine antique-gold ornament accents, warm golden-hour light, "
         "rich jewel tones (saffron, indigo, vermilion, emerald, ivory, antique gold), delicate brushwork, elegant composition, "
         "crisp readable silhouette, culturally respectful, welcoming to players worldwide, masterpiece, highly detailed, original design")
ICON = ("single object centred, isolated on a plain transparent background, three-quarter view, soft painted shading, "
        "thin antique-gold outline, readable at 64 px, no background scenery")
SCENE = "cinematic wide composition, layered depth, soft atmospheric haze, festival sky full of colourful paper kites"
NEG = ("text, letters, numbers, watermark, logo, brand name, signature, frame text, blurry, low resolution, jpeg artifacts, "
       "photograph, plastic 3D render, oversaturated neon, extra fingers, deformed hands, distorted face, "
       "glass-coated string, metal wire string, nylon string, blades, blood, weapons, real-world flags, copyrighted characters")
NEG_ICON = NEG + ", background scenery, cropped object, multiple objects, drop shadow box"

rows = []  # (category, file, size, fmt, where, priority, note, prompt, neg)
def add(cat, file, size, fmt, where, pri, note, prompt, neg=NEG):
    rows.append((cat, file, size, fmt, where, pri, note, prompt, neg))
def icon(cat, file, where, note, subject, size="512x512", pri="High"):
    add(cat, file, size, "PNG (transparent)", where, pri, note, f"{subject}, {ICON}, {STYLE}", NEG_ICON)

# ---------- Branding / store ----------
add("Branding", "app_icon.png", "1024x1024 (export 512x512 for Play Store)", "PNG (no transparency)", "Android launcher + Google Play listing", "Critical",
    "Must read at 48 px. One bold kite, no text. Keep important art inside the centre 66 % (Android masks corners).",
    f"A single majestic diamond-shaped paper kite in saffron and vermilion with an ornate gold-leaf mandala pattern, soaring against a glowing sunset sky, flowing ribbon tail curling around it, {STYLE}, bold simple icon composition, high contrast")
add("Branding", "icon_adaptive_foreground.png", "432x432", "PNG (transparent)", "Android adaptive icon foreground layer", "Critical",
    "Same kite as app_icon, alone, inside the centre 288x288 safe zone. Transparent around it.",
    f"The same ornate saffron and vermilion diamond paper kite with curling ribbon tail, centred, {ICON}, {STYLE}", NEG_ICON)
add("Branding", "icon_adaptive_background.png", "432x432", "PNG", "Android adaptive icon background layer", "Critical",
    "Only a soft sky gradient with faint gold pattern. No objects.",
    f"Soft warm sunset sky gradient from deep indigo to saffron with very faint gold-leaf geometric jali pattern, seamless, no objects, {STYLE}")
add("Branding", "logo_kitefall.png", "2048x1024", "PNG (transparent)", "Title screen, splash, store art", "Critical",
    "AI cannot write text well – generate the ORNAMENT ONLY, then the designer adds the word 'KITEFALL' by hand in a custom calligraphic font.",
    f"Ornamental title cartouche emblem with empty centre space for a word, gold-leaf arabesque scrollwork, a small paper kite and ribbon flourish on the left, symmetrical, isolated on transparent background, {STYLE}", NEG_ICON)
add("Branding", "splash_screen.png", "2400x1080 (landscape)", "PNG/JPG", "Unity splash / first loading screen", "Critical",
    "Leave the centre-top clear for the logo. Keep key art in the middle 1920x1080 so 16:9 and 20:9 phones both look good.",
    f"Breathtaking panoramic view of an old walled city at golden hour, dense stacked brick rooftops with carved wooden balconies, hundreds of colourful paper kites filling the sky, a young flyer silhouetted on a rooftop holding a wooden spool, {SCENE}, {STYLE}")
for i, (name, desc) in enumerate([
        ("loading_01_oldcity.jpg", "children and families on old city rooftops at sunrise preparing kites, marigold garlands and string lights"),
        ("loading_02_coast.jpg", "seaside cliff houses painted white and teal, kites flying over a turquoise sea with fishing boats"),
        ("loading_03_river.jpg", "a riverside spring festival with lantern boats, temple steps and kites reflected in calm water"),
        ("loading_04_royal.jpg", "a grand palace garden with marble pavilions, fountains and royal kite flyers on terraces"),
        ("loading_05_mountain.jpg", "a terraced mountain village with prayer-flag-like coloured cloth and kites over snowy peaks"),
        ("loading_06_storm.jpg", "an island archipelago under dramatic storm clouds with brave kites cutting through the wind"),
        ], 1):
    add("Loading screens", name, "2400x1080 (landscape)", "JPG (quality 85)", "Random loading screen behind the progress bar + tips", "High",
        "Bottom 20 % will be covered by the loading bar and tip text – keep it calm (sky or ground).",
        f"{desc}, {SCENE}, {STYLE}")
add("Store listing", "play_feature_graphic.png", "1024x500", "PNG/JPG (no transparency)", "Google Play feature graphic", "Critical",
    "No text in the image (Play adds its own). Two rival kites crossing strings in the centre.",
    f"Two ornate paper kites, one saffron and one indigo, crossing strings dramatically in a glowing festival sky above old city rooftops, sense of a friendly rooftop duel, {SCENE}, {STYLE}")
for i in range(1, 9):
    add("Store listing", f"store_screenshot_frame_{i:02d}.png", "1920x1080", "PNG", "Google Play screenshots (decorative frame/background)", "Medium",
        "Real gameplay screenshots go inside. This is only the decorative ornamental frame + background band for marketing text.",
        f"Decorative landscape banner frame with an empty centre window, gold-leaf border with floral arabesques and tiny kites in the corners, deep indigo velvet background, {STYLE}", NEG_ICON)

# ---------- Menus ----------
for f, w, d in [
    ("bg_home.jpg", "Home screen background", "a cosy rooftop at golden hour with a wooden day-bed, potted plants, wooden spools and a kite rack, city of rooftops behind, kites in the sky, gentle and inviting"),
    ("bg_shop.jpg", "Shop background", "an old bazaar kite stall with kites hanging from wooden beams, spools of coloured cotton thread on shelves, warm lantern light"),
    ("bg_hangar.jpg", "Hangar / loadout background", "a craftsman's kite workshop with bamboo sticks, tissue paper sheets, glue pots, painted wooden spools, soft window light"),
    ("bg_wardrobe.jpg", "Wardrobe / character screen background", "an elegant tailor's room with folded festive clothes, carved wooden mirror frame, embroidered curtains, a soft spotlight circle on the floor in the centre"),
    ("bg_squad_lobby.jpg", "Squad lobby background", "a wide rooftop terrace at dusk with four empty spots marked by rugs for friends to stand, string lights overhead"),
    ("bg_pass.jpg", "Season pass background", "a grand festive night sky with fireworks-like lantern lights, golden kites rising in a spiral, royal ornamental feel"),
    ("bg_results_win.jpg", "Victory results screen", "a triumphant golden sunrise sky with confetti of paper petals and a single proud kite rising high"),
    ("bg_results_lose.jpg", "Defeat results screen", "a calm twilight sky with a lone kite drifting gently down, peaceful and hopeful, not sad"),
    ("bg_settings.jpg", "Settings / profile background", "a soft blurred rooftop evening with very low detail, subtle gold pattern, calm"),
]:
    add("Menu backgrounds", f, "2400x1080 (landscape)", "JPG (quality 85)", w, "High",
        "UI panels sit on top – keep the centre area soft/low-contrast. Will be dimmed 20-30 % in game.",
        f"{d}, {SCENE if 'kite' in d else 'cinematic composition, soft depth of field'}, {STYLE}")

# ---------- Kites ----------
KITES = {
    "guddi": "classic South Asian diamond fighter kite (guddi/patang) with a curved bamboo spine and bow, small triangular tail flap",
    "pipa": "small square fighter kite (pipa) with a central bamboo spine and arched cross-bow, light and nimble",
    "rokkaku": "six-sided tall hexagonal kite (rokkaku style) with two cross spars, bold painted panels",
    "pakpao": "small Thai-style diamond fighter kite (pakpao) with long thin tail strings, agile shape",
    "gudiparan": "Afghan-style fighter kite (gudiparan) with a wide diamond body and a tiny tassel tail, elegant curved bow",
    "chula": "large Thai star-shaped kite (chula) with five pointed bamboo arms, powerful and grand",
    "wau": "Malaysian moon kite (wau bulan) with a crescent tail, wide curved wings and intricate floral cut-paper decoration",
}
for k, d in KITES.items():
    icon("Kites", f"kite_{k}_card.png", "Hangar kite cards, shop, loadout picker",
         f"Main card art for the {k} kite class. Plain ivory paper colour so the in-game skin colour stands out. Must match the in-game 3D shape.",
         f"An elegant {d}, made of translucent tissue paper on bamboo sticks, ivory and gold base colours with subtle hand-painted border lines", "1024x1024", "Critical")
PAL = [("Sunset", "vermilion red and warm butter yellow", "#F25A3C/#FFE08A"),
       ("Monsoon", "royal blue and pale sky blue", "#2F6FDB/#9FD8FF"),
       ("Basant", "mustard yellow and fresh leaf green", "#F5C400/#3BA55C"),
       ("Night Bazaar", "deep violet and rose pink", "#3A2A6B/#FF7AC8"),
       ("Sea Glass", "teal and mint white", "#21A6A6/#E8FFF6"),
       ("Saffron", "saffron orange and pure white", "#FF8C1A/#FFFFFF"),
       ("Jade", "jade green and soft lime", "#1F8F5F/#D9F2B4"),
       ("Ember", "deep crimson and charcoal black", "#B22222/#1A1A1A")]
PATTERNS = ["Mughal floral border", "geometric jali lattice", "sunburst rays", "paisley motifs", "lotus mandala", "block-print stripes"]
for ki, k in enumerate(KITES):
    for i, (p, colours, hexes) in enumerate(PAL):
        rarity = "Common" if i < 3 else "Rare" if i < 6 else "Epic" if i == 6 else "Legendary (pass only)"
        pat = PATTERNS[(i + ki) % 6]
        add("Kite skins", f"skin_{k}_{i}_{p.lower().replace(' ', '_')}.png", "1024x1024", "PNG (tileable not needed)", f"Kite paper texture for {k} · {p} ({rarity}); also shop thumbnail", "High" if i < 3 else "Medium",
            f"FLAT front-view texture of the kite sail only (no sticks, no sky), filling the whole square – it is UV-mapped onto the 3D kite. Colours {hexes}. Pattern: {pat}. Higher rarity = more gold detail.",
            f"Flat front-on artwork of a {k} kite sail as a design sheet, {colours} colour scheme, {pat} hand-painted on thin tissue paper, {'rich gold-leaf ornament' if i >= 6 else 'fine ink linework'}, symmetrical, fills the frame edge to edge, {STYLE}",
            NEG + ", sky, background, perspective, bamboo sticks, hands")

# ---------- Loadout items ----------
TAILS = {"none": "an empty dashed circle meaning no tail", "short_ribbon": "a short single silk-like ribbon tail curling in the wind",
         "double_ribbon": "two parallel ribbon tails fluttering", "long_tail": "a very long flowing cloth tail with small bows",
         "tassel": "a decorative cotton tassel with tiny beads"}
for k, d in TAILS.items():
    icon("Dum (tails)", f"dum_{k}.png", "Loadout picker, hangar", "Tail item icon. Colours neutral ivory/gold so it works with every kite.", d)
DORS = {"sooti": "a hand-wound ball of plain natural cotton thread, honest and simple",
        "reshmi": "a smooth lustrous ball of fine silk-like thread with soft sheen",
        "tez_dhaar": "a tightly wound ball of thick deep-crimson dyed cotton thread with a bold aggressive look",
        "mazboot": "a heavy thick braided cotton cord wound in a ball, sturdy and strong",
        "lambi": "a very large ball of long thin cotton thread, generous and abundant",
        "lachakdaar": "a springy loosely coiled ball of soft flexible cotton thread"}
for k, d in DORS.items():
    icon("Dor (string)", f"dor_{k}.png", "Loadout picker, hangar, shop",
         "SAFETY: show only natural cotton/silk-like thread. Never glass coating, metal or nylon. Stats feeling comes from colour & thickness only.", d)
CHARKHIS = {"lakri": "a classic carved wooden kite spool (charkhi) with two handles, warm walnut wood",
            "speed": "a slim lightweight wooden kite spool with streamlined spokes and red lacquer stripes",
            "bari": "a large wide-drum wooden kite spool holding lots of thread",
            "brake": "a sturdy wooden kite spool with a leather friction strap wrapped around one handle",
            "gear": "an ornate wooden kite spool with carved cog-like rims and brass-coloured painted details"}
for k, d in CHARKHIS.items():
    icon("Charkhi (spool)", f"charkhi_{k}.png", "Loadout picker, hangar, shop", "Spool item icon. Wooden look, painted details allowed.", d)
TUNINGS = {"waxed_coat": "a small pot of beeswax with a cotton thread passing over it", "double_twist": "two cotton threads twisting together into one",
           "light_thread": "a feather-light wisp of thread floating upward", "oiled_axle": "a tiny oil can dripping one drop onto a wooden axle",
           "wide_drum": "a wide wooden drum spool seen from the side", "soft_grip": "a soft wrapped cloth handle grip"}
for k, d in TUNINGS.items():
    icon("Tuning", f"tuning_{k}.png", "Hangar tuning slots", "Small upgrade badge. Round medallion shape with the object inside.",
         f"a round ornate gold medallion badge containing {d}", "256x256", "Medium")

# ---------- Characters ----------
CHARS = [("rookie_m", "Ali", "a cheerful young man from an old South Asian city, short dark hair, warm brown skin, street kurta"),
         ("rookie_f", "Sana", "a confident young woman with long braided dark hair, light brown skin, bright yellow and green festive outfit"),
         ("zara", "Zara", "a stylish young woman from Lahore with a high bun, deep violet and orange night-bazaar outfit, playful smile"),
         ("rafay", "Rafay", "a tall broad-shouldered young man from Karachi, buzz-cut hair, coastal linen shirt in white and teal"),
         ("lucas", "Lucas", "a friendly young man from Lisbon with wavy chestnut hair, fair skin, blue monsoon rain jacket"),
         ("aiko", "Aiko", "an energetic young woman from Osaka with a sleek black bob, pink and yellow festival-lights jacket"),
         ("niran", "Niran", "a calm young man from Chiang Mai with short black hair, navy rooftop hoodie with orange trim"),
         ("farid", "Farid", "a rugged young man from Kabul with tousled dark hair and short beard, red mountain parka"),
         ("siti", "Siti", "a quick athletic young woman from Kuala Lumpur with a long ponytail, green and lime storm-runner outfit"),
         ("arjun", "Arjun", "a dignified young man from Ahmedabad with neat black hair, cream and gold royal sherwani")]
for cid, name, d in CHARS:
    add("Characters", f"char_{cid}_portrait.png", "1024x1024", "PNG (transparent)", f"{name} – portrait in wardrobe, party lobby, profile, name tags", "Critical",
        f"Head-and-shoulders, friendly expression, looking slightly left. Must match the in-game 3D look (skin/hair colours from CharacterCatalog). Original person – not a real celebrity. Age 18+ look.",
        f"Portrait of {name}, {d}, holding a small paper kite near the shoulder, warm smile, three-quarter view, isolated on transparent background, {STYLE}", NEG_ICON)
    add("Characters", f"char_{cid}_fullbody.png", "1024x1536 (portrait)", "PNG (transparent)", f"{name} – character select card, shop card", "High",
        "Full body standing pose, feet visible, holding a wooden kite spool. Same face & outfit as the portrait.",
        f"Full-body character art of {name}, {d}, standing proudly on tiptoe holding a carved wooden kite spool, kite string rising out of frame, isolated on transparent background, {STYLE}", NEG_ICON)
OUTFITS = [("kurta_street", "Street Kurta", "an off-white cotton kurta with slate grey trousers"),
           ("basant_yellow", "Basant Yellow", "a mustard yellow kurta with green embroidered borders"),
           ("rooftop_hoodie", "Rooftop Hoodie", "a navy hoodie with orange trim and loose joggers"),
           ("oldcity_waistcoat", "Old City Waistcoat", "an ivory kurta under a maroon embroidered waistcoat"),
           ("monsoon_jacket", "Monsoon Jacket", "a royal blue light rain jacket over dark trousers"),
           ("festival_lights", "Festival Lights", "a pink jacket with tiny yellow embroidered light motifs"),
           ("coastal_linen", "Coastal Linen", "a mint-white linen shirt with teal sash"),
           ("royal_sherwani", "Royal Sherwani", "a cream sherwani with antique gold embroidery"),
           ("mountain_parka", "Mountain Parka", "a crimson padded parka with charcoal trims"),
           ("storm_runner", "Storm Runner", "a jade green athletic wind jacket with lime panels"),
           ("night_bazaar", "Night Bazaar", "a deep violet kurta with glowing orange mirror-work embroidery"),
           ("champion_kit", "Champion Kit", "a black and gold champion outfit with gold kite emblem embroidery")]
for oid, name, d in OUTFITS:
    add("Outfits", f"outfit_{oid}.png", "768x1024", "PNG (transparent)", f"Wardrobe + shop card for '{name}'", "High",
        "Show the outfit on a faceless neutral mannequin/dress form so it works for every character.",
        f"{d} displayed on an elegant faceless wooden tailor's mannequin, front view, isolated on transparent background, {STYLE}", NEG_ICON)
ACCS = [("acc_topi", "a white crocheted prayer cap"), ("acc_cap", "an orange-red baseball-style cap"), ("acc_turban", "an orange festival turban with a gold brooch"),
        ("acc_headband", "a blue cotton sweatband"), ("acc_shades", "stylish black sunglasses"), ("acc_goggles", "teal aviator wind goggles with leather strap"),
        ("acc_kitebag", "a brown leather kite carrying bag"), ("acc_scarf", "a long crimson woven scarf"), ("acc_spare_kite", "a spare yellow paper kite tied with a strap"),
        ("acc_pouch", "a small brown leather belt pouch"), ("acc_sash", "a green embroidered festival sash"), ("acc_holster", "a leather holster holding a small wooden spool")]
for aid, d in ACCS:
    icon("Accessories", f"{aid}.png", "Wardrobe accessory slots, shop", "Accessory icon. Same angle and lighting for all 12.", d)
EMOTES = [("bo_kata", "a joyful raised fist with a cut kite string floating away (the famous 'Bo Kata!' cheer)"),
          ("thumbs_up", "a thumbs-up hand"), ("taunt_wave", "a friendly waving hand"), ("victory_dance", "a small dancing figure silhouette with motion swirls"),
          ("clap", "two clapping hands with sparkle lines"), ("bow", "a figure bowing respectfully"), ("spin", "a figure spinning with circular motion lines"),
          ("laugh", "a laughing face with closed happy eyes"), ("shrug", "a shrugging figure"), ("dhol_beat", "a festive double-headed dhol drum with sticks"),
          ("salute", "a hand saluting"), ("facepalm", "a hand on a forehead in a facepalm")]
for eid, d in EMOTES:
    icon("Emotes", f"emote_{eid}.png", "Emote wheel, wardrobe, shop", "Round emote button art. Expressive and fun, readable at 64 px.",
         f"a round ornate medallion with {d}", "256x256")

# ---------- Arenas & modes ----------
ARENAS = [("old_city", "Old City Rooftops", "dense old walled city rooftops with brick houses, carved wooden balconies, water tanks and laundry lines"),
          ("coastal_cliffs", "Coastal Cliffs", "white and teal houses on cliffs above a sparkling sea, lighthouse in the distance"),
          ("river_festival", "River Festival", "riverside houses with festival lanterns, boats and stone ghats"),
          ("royal_field", "Royal Field", "a vast palace ground with domed pavilions and manicured gardens"),
          ("mountain_valley", "Mountain Valley", "a terraced village in a green valley with snowy peaks behind"),
          ("storm_archipelago", "Storm Archipelago", "small rocky islands with stilt houses under dramatic dark storm clouds and shafts of light")]
for aid, name, d in ARENAS:
    add("Arenas", f"arena_{aid}_card.jpg", "1280x720", "JPG (quality 90)", f"Arena select card + match intro for '{name}'", "Critical",
        "Arena name is added by UI – no text in image. Must look like the in-game arena layout (rooftops in the foreground).",
        f"{d}, kites flying overhead, {SCENE}, {STYLE}")
    add("Arenas", f"arena_{aid}_minimap.png", "512x512", "PNG", f"Map camera / roof select overlay for '{name}'", "Medium",
        "Top-down stylised map like an old painted city map. Designer should trace the real roof layout from a top-down screenshot.",
        f"Top-down illustrated antique map of {d}, painted like an old manuscript city plan, soft colours, {STYLE}")
MODES = [("quick_duel", "Quick Duel", "two kites crossing strings one-on-one"), ("ranked_duel", "Ranked Duel", "two kites duelling with a golden laurel wreath behind"),
         ("ffa", "Free-for-All", "eight colourful kites tangling in a crowded sky"), ("private_room", "Private Room", "a carved wooden door slightly open to a rooftop with kites"),
         ("team_battle", "Team Battle", "two groups of kites, saffron team and indigo team, facing each other"),
         ("last_kite_flying", "Last Kite Flying", "one proud kite remaining high above many falling kites"),
         ("training", "Training", "a kite tied to a practice pole on a quiet rooftop with a friendly helper"),
         ("tutorial", "Tutorial", "an elder teaching a child to fly a kite on a rooftop")]
for mid, name, d in MODES:
    add("Game modes", f"mode_{mid}_card.jpg", "1024x576", "JPG (quality 90)", f"Play menu mode card '{name}'", "Critical",
        "Card frame & title are UI. Keep the left third calmer for the mode name text.",
        f"{d}, {SCENE}, {STYLE}")
TRAITS = [("high", "a tall rooftop with an upward arrow"), ("sheltered", "a rooftop with a wall blocking wind swirls"),
          ("upwind", "a wind gust swirl with an arrow"), ("crate_spot", "a small wooden supply crate with a kite symbol"),
          ("edge", "a rooftop edge with a caution stripe"), ("cables", "tangled overhead cables (warning)")]
for tid, d in TRAITS:
    add("Roof traits", f"trait_{tid}.png", "256x256", "PNG (transparent)", "Roof select screen, map labels", "Medium", "WHITE glyph on transparent background (the game tints it). Must read at 40 px.", f"a simple flat white silhouette symbol of {d}, pure white on transparent background, clean edges, no shading, game UI glyph", NEG_ICON + ", colours, gradients")
for lid, d in [("plank", "a wooden plank bridge between rooftops"), ("rope_bridge", "a swaying rope bridge"), ("ladder", "a wooden ladder"), ("stairs", "stone stairs")]:
    icon("Roof traits", f"link_{lid}.png", "Roof hopping path markers", "Path icon shown on the map for roof hopping.", f"a small painted icon of {d}", "256x256", "Low")

# ---------- Economy ----------
for cid, d, note in [("coins", "a stack of three antique gold coins with a tiny kite stamped on them", "Soft currency. Gold."),
                     ("gems", "a cluster of faceted ruby and emerald gems", "Premium currency. Must look clearly different from coins."),
                     ("parts", "a bundle of bamboo sticks, a paper sheet and a small thread spool tied together", "Crafting parts.")]:
    icon("Currency", f"currency_{cid}.png", "Top bar, shop, rewards", note + " Also export 64x64 version.", d, "512x512", "Critical")
for sid, d in [("gems_small", "a small velvet pouch of gems"), ("gems_medium", "a larger cloth bag overflowing with gems"), ("gems_large", "an ornate wooden treasure chest filled with gems")]:
    icon("Shop", f"shop_{sid}.png", "Gem store tiles", "Bigger pack = bigger, richer art.", d, "512x512", "High")
icon("Shop", "shop_no_ads.png", "Store tile 'No ads'", "Calm, premium. A peaceful clear sky symbol.", "a peaceful clear sky medallion with a single kite and a soft gold ribbon seal", "512x512", "High")
icon("Shop", "shop_premium_pass.png", "Store tile 'Premium Season Pass'", "Royal look – gold and indigo.", "a royal golden scroll pass with an indigo ribbon and a kite seal", "512x512", "High")
add("Season pass", "pass_banner.jpg", "2048x512", "JPG", "Season pass header banner", "High",
    "Wide banner, season title added by UI. Change art each season.",
    f"Wide ornamental festival banner, a spiral of golden kites rising into a twilight sky, gold-leaf arabesque borders top and bottom, {STYLE}")
for t, d in [("free_track", "a simple bronze ribbon medallion"), ("premium_track", "a jewelled golden crown medallion"), ("level_node", "a small round gold coin-like level token")]:
    icon("Season pass", f"pass_{t}.png", "Season pass tracks", "Pass UI piece.", d, "256x256", "Medium")
for cid, d in [("better_dor", "a wooden crate with a cotton thread ball symbol"), ("better_charkhi", "a wooden crate with a spool symbol"),
               ("patch_kit", "a small cloth kit with paper patches and glue pot"), ("dor_refill", "a crate with fresh thread spools")]:
    icon("Sky crates", f"crate_{cid}.png", "In-match sky crate pickup + HUD", "Must read clearly from far away in the sky. Small parachute-like cloth is fine.", d, "512x512", "High")
icon("Rewards", "daily_chest.png", "Daily reward chest", "Closed version. Also make 'daily_chest_open.png' with glow.", "an ornate carved wooden chest with brass corners and a kite emblem", "512x512", "High")
icon("Rewards", "daily_chest_open.png", "Daily reward chest opened", "Same chest open with golden light.", "the same ornate carved wooden chest open with warm golden light and sparkles rising", "512x512", "High")

# ---------- Rank ----------
for rid, d in [("bronze", "bronze"), ("silver", "polished silver"), ("gold", "rich gold"), ("platinum", "pale platinum with pearl"),
               ("diamond", "sparkling diamond-blue crystal"), ("master", "royal purple and gold"), ("sky_legend", "radiant gold with sun rays and a flying kite crest")]:
    icon("Rank badges", f"rank_{rid}.png", "Ranked screen, profile, results", "Shield/emblem. Each tier more ornate than the last. Same shape family.",
         f"a heraldic shield emblem in {d} with a diamond kite crest and laurel leaves", "512x512", "High")

# ---------- Cosmetic effects ----------
for i, n in enumerate(["Ribbon sparkle", "Rainbow trail", "Firefly tail", "Petal drift"]):
    icon("Effects", f"dumfx_{i}_{n.lower().replace(' ', '_')}.png", "Shop/wardrobe preview icon for tail effect", "Preview only – real effect is a particle system.", f"a kite tail with a {n.lower()} effect", "512x512", "Medium")
for i, n in enumerate(["Paper confetti", "Lantern burst", "Feather swirl", "Firework"]):
    icon("Effects", f"cutfx_{i}_{n.lower().replace(' ', '_')}.png", "Shop preview icon for cut effect", "Preview only.", f"a burst of {n.lower()} in the sky", "512x512", "Medium")
for i, n in enumerate(["Indigo dye", "Rose dye", "Gold thread look", "Glow (visual only)"]):
    icon("Effects", f"dorskin_{i}.png", "Shop preview for dor skin", "Cotton thread look only (safety rule).", f"a ball of cotton thread in {n.lower()} colour", "512x512", "Medium")
for i, n in enumerate(["Painted wood", "Brass rim", "Truck art", "Mirror work"]):
    icon("Effects", f"charkhiskin_{i}.png", "Shop preview for charkhi skin", "Same spool shape, only decoration changes.", f"a wooden kite spool decorated with {n.lower()} style ornament", "512x512", "Medium")

# ---------- Social / HUD ----------
for q, d in [("up", "an upward arrow"), ("help", "a raised hand asking for help"), ("attack", "two crossing kite strings"), ("defend", "a small shield with a kite"),
             ("crate", "a wooden crate"), ("cut", "a cut thread with a spark"), ("thanks", "a hand on heart"), ("gg", "two hands shaking")]:
    add("Quick chat", f"chat_{q}.png", "256x256", "PNG (transparent)", "Quick chat wheel", "Medium", "WHITE glyph on transparent background (the game tints it). Readable at 48 px.", f"a simple flat white silhouette symbol of {d}, pure white on transparent background, clean edges, no shading, game UI glyph", NEG_ICON + ", colours, gradients")
for h, d in [("dheel", "an unwinding spool with thread flowing out (release)"), ("kheench", "a hand pulling thread back hard (pull)"),
             ("mic", "a vintage microphone"), ("speaker", "a vintage speaker horn"), ("map", "a folded old paper map"), ("pack_up", "a kite folded under an arm"),
             ("launch", "a kite rising from a hand"), ("wind", "wind swirl lines"), ("snap_warning", "a straining thread about to snap"), ("emote", "a smiling face")]:
    add("HUD", f"hud_{h}.png", "256x256", "PNG (transparent)", "In-match HUD buttons", "High", "WHITE glyph on transparent background (the game tints it). Simple bold shape, thumb-size button.", f"a bold simple flat white silhouette icon of {d}, pure white on transparent background, clean vector-like edges, no shading, game UI glyph", NEG_ICON + ", colours, gradients")
for m in ["play3", "win2", "cut5", "ffa1", "sweet", "nosnap", "assist", "duel3"]:
    d = {"play3": "three small kites in a row", "win2": "a trophy with two stars", "cut5": "five cut thread pieces", "ffa1": "a crowded sky of kites",
         "sweet": "a glowing sweet-spot target", "nosnap": "an unbroken thread loop", "assist": "two hands passing a kite", "duel3": "crossed kite strings"}[m]
    icon("Missions", f"mission_{m}.png", "Daily missions list", "Mission icon, round.", f"a round mission badge with {d}", "256x256", "Medium")

# ---------- Onboarding ----------
for i, (f, d) in enumerate([("onboard_01_welcome.png", "a warm rooftop welcome scene with friends waving and kites rising"),
                            ("onboard_02_fly.png", "a hand holding a spool and the kite climbing in the sky with motion lines"),
                            ("onboard_03_cut.png", "two kites crossing strings with a dramatic spark of friendly competition"),
                            ("onboard_04_roof.png", "players hopping between rooftops on a plank bridge"),
                            ("onboard_05_squad.png", "four friends on one rooftop celebrating together")], 1):
    add("Onboarding", f, "1600x900", "PNG", "Tutorial / first launch slides", "High", "Simple story illustration, room at the bottom for one line of text.", f"{d}, {STYLE}")

# ---------- Environment textures ----------
for f, d in [("tex_brick_wall.png", "old weathered red brick wall"), ("tex_plaster_wall.png", "lime-washed plaster wall with soft stains"),
             ("tex_roof_floor.png", "rooftop floor of terracotta tiles"), ("tex_wood_balcony.png", "carved dark wooden balcony panel (jharokha)"),
             ("tex_jali_pattern.png", "stone jali lattice screen pattern"), ("tex_fabric_rug.png", "woven rug pattern")]:
    add("Environment textures", f, "1024x1024 (tileable)", "PNG", "Optional upgrade for the procedural ArenaBuilder textures", "Low",
        "MUST tile seamlessly. Flat lighting, no shadows baked in. Code currently draws procedural textures – these replace them later.",
        f"Seamless tileable texture of {d}, flat even lighting, top-down orthographic, stylised painterly, game texture, {STYLE}", NEG + ", perspective, shadows, objects")

wb = Workbook()
hdr_fill = PatternFill("solid", fgColor="3A2A6B")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D9CFB8")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

# Style guide
sg = wb.active
sg.title = "Style Guide"
sg.column_dimensions["A"].width = 28
sg.column_dimensions["B"].width = 120
sg["A1"] = "KITEFALL – Image Art Style Guide"
sg["A1"].font = Font(bold=True, size=18, color="3A2A6B")
guide = [
    ("Art direction", "INTERNATIONAL game: classical hand-painted storybook look that welcomes players worldwide + premium mobile polish. Gouache feel, fine gold ornament, warm golden-hour light. Regional items (an arena, a character, a kite) may show their own culture respectfully; menus, branding and UI stay globally neutral. Elegant, festive, friendly – never dark or violent."),
    ("Master style prompt (add to every prompt)", STYLE),
    ("Icon add-on", ICON),
    ("Scene add-on", SCENE),
    ("Master negative prompt", NEG),
    ("Colour palette", "Saffron #FF8C1A · Vermilion #F25A3C · Indigo #3A2A6B · Royal blue #2F6FDB · Emerald/Jade #1F8F5F · Mustard #F5C400 · Ivory #F4EFE6 · Antique gold #B08A2E · Teal #21A6A6"),
    ("Consistency", "Generate all items of one category in the same session with the same seed/style reference. Use one approved image as the style reference (image prompt / --sref / IP-Adapter) for the rest."),
    ("Text in images", "NEVER let AI draw text. All words (logo name, mode names, prices) are added by the designer or by the game UI."),
    ("SAFETY rule (kite string)", "Kite string (dor) must look like natural cotton or silk-like thread only. Never show glass-coated, metal or nylon string, blades or injuries."),
    ("Originality", "All art must be original. No brand logos, no real celebrities, no characters from other games/films, no copied photos. Characters are fictional adults (18+ look)."),
    ("Icons", "Transparent PNG, object centred with ~8 % padding, same light direction (top-left) for every icon. Check readability by shrinking to 64 px."),
    ("Kite skins", "Flat front-on sail texture filling the square – it is UV-mapped onto the 3D kite. No sky, no sticks, no perspective."),
    ("Export & import in Unity", "Put files in Assets/Art/<Category>/. Icons: Texture Type = Sprite (2D and UI), Max Size 512, compression ASTC 6x6. Backgrounds: Max Size 2048, ASTC 8x8. Kite skins: Texture Type = Default, Max Size 1024, mipmaps ON."),
    ("Suggested tools", "Midjourney / Leonardo / Stable Diffusion XL / Firefly / DALL·E. Midjourney tip: add '--ar 16:9' for backgrounds, '--ar 1:1' for icons, '--style raw --s 250'. Then upscale and clean up in Photoshop/Krita."),
    ("Different file names?", "No problem. In Unity open Kitefall > Art > Image Mapper. Drag each image onto its slot, or press Auto-match. Or type your own file name in the last column of 'Image List', save this sheet as CSV and press 'Import CSV' in the mapper. Files named only with the # number (e.g. 45.png) can be matched with 'Match by # number'."),
    ("Licence check", "Use an AI tool whose terms allow commercial use, and keep the prompt + date for every final image (record it in the 'Status / notes' column)."),
]
r = 3
for k, v in guide:
    sg.cell(row=r, column=1, value=k).font = Font(bold=True, color="7A2E2E")
    c = sg.cell(row=r, column=2, value=v)
    c.alignment = wrap
    sg.cell(row=r, column=1).alignment = wrap
    sg.row_dimensions[r].height = max(20, 15 * (len(v) // 110 + 1))
    r += 1

# Image list
ws = wb.create_sheet("Image List")
headers = ["#", "Category", "File name", "Size (px)", "Format", "Where used in game", "Priority", "Designer note", "AI prompt", "Negative prompt", "Status", "Status / notes", "Your file name (fill if different)"]
widths = [5, 18, 34, 22, 18, 34, 10, 50, 90, 50, 12, 30, 34]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=i, value=h)
    c.fill = hdr_fill; c.font = hdr_font; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = border
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 28
pri_colour = {"Critical": "F8CBAD", "High": "FFE699", "Medium": "E2EFDA", "Low": "DDEBF7"}
band = PatternFill("solid", fgColor="FBF8F1")
for n, row in enumerate(rows, 1):
    vals = [n] + list(row) + ["To do", "", ""]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=n + 1, column=j, value=v)
        c.alignment = wrap; c.border = border
        if n % 2 == 0: c.fill = band
    ws.cell(row=n + 1, column=7).fill = PatternFill("solid", fgColor=pri_colour.get(row[5].split()[0], "FFFFFF"))
    ws.row_dimensions[n + 1].height = 120
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
dv = DataValidation(type="list", formula1='"To do,In progress,Review,Approved,In Unity"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"K2:K{len(rows) + 1}")

# Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Category"; sm["B1"] = "Images"; sm["C1"] = "Critical"; sm["D1"] = "Done (Approved / In Unity)"
for c in ("A1", "B1", "C1", "D1"):
    sm[c].fill = hdr_fill; sm[c].font = hdr_font
cats = []
for row in rows:
    if row[0] not in cats: cats.append(row[0])
last = len(rows) + 1
for i, cat in enumerate(cats, 2):
    sm.cell(row=i, column=1, value=cat)
    sm.cell(row=i, column=2, value=f"=COUNTIF('Image List'!B2:B{last},A{i})")
    sm.cell(row=i, column=3, value=f"=COUNTIFS('Image List'!B2:B{last},A{i},'Image List'!G2:G{last},\"Critical\")")
    sm.cell(row=i, column=4, value=f"=COUNTIFS('Image List'!B2:B{last},A{i},'Image List'!K2:K{last},\"Approved\")+COUNTIFS('Image List'!B2:B{last},A{i},'Image List'!K2:K{last},\"In Unity\")")
t = len(cats) + 2
sm.cell(row=t, column=1, value="TOTAL").font = Font(bold=True)
for col in "BCD":
    sm[f"{col}{t}"] = f"=SUM({col}2:{col}{t - 1})"; sm[f"{col}{t}"].font = Font(bold=True)
sm.column_dimensions["A"].width = 26
for col in "BCD": sm.column_dimensions[col].width = 16
wb.move_sheet("Summary", offset=-1)
wb.save(OUT)
print(len(rows), "rows,", len(cats), "categories")

BUILD_ONLY = {"Branding", "Store listing"}
cs = ["// <auto-generated> by Tools/ArtList/make_assets.py from docs/Kitefall_Image_Assets.xlsx – do not edit by hand.",
      "namespace Kitefall.UI", "{", "    public static partial class ArtSlots", "    {", "        public static readonly ArtSlot[] All =", "        {"]
for n, row in enumerate(rows, 1):
    stem = row[1].rsplit(".", 1)[0]
    size = row[2].split(" ")[0]
    cs.append(f'            new ArtSlot({n}, "{stem}", "{row[0]}", "{size}", {"true" if row[0] == "Store listing" or stem.startswith(("app_icon", "icon_adaptive")) else "false"}),')
cs += ["        };", "    }", "}", ""]
open(sys.argv[2], "w").write("\n".join(cs))

# Editor-only keyword table: what each image shows (prompt subject without the shared style text) – used to match
# artist files whose names describe the picture ("Golden Kite Bazaar at Sunset") instead of the slot id.
def subject(prompt):
    prompt = prompt.replace("fills the frame edge to edge", "").replace("out of frame", "")
    for part in (STYLE, ICON, SCENE):
        prompt = prompt.replace(part, "")
    return re.sub(r"[^A-Za-z ]+", " ", prompt)
kw = ["// <auto-generated> by Tools/ArtList/make_assets.py – do not edit by hand.", "using System.Collections.Generic;", "",
      "namespace Kitefall.EditorTools", "{", "    internal static class ArtSlotKeywords", "    {",
      "        public static readonly Dictionary<string, string> ById = new Dictionary<string, string>", "        {"]
for row in rows:
    stem = row[1].rsplit(".", 1)[0]
    # Kite skins: only colours + pattern (their palette names like "Night Bazaar" would steal scene matches).
    extra = "" if row[0] == "Kite skins" else " " + row[4] + " " + stem.replace("_", " ")
    words = " ".join((row[0] + extra + " " + subject(row[7])).lower().split())
    kw.append(f'            {{ "{stem}", "{words}" }},')
kw += ["        };", "    }", "}", ""]
if len(sys.argv) > 3:
    open(sys.argv[3], "w").write("\n".join(kw))
